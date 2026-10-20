"""
Writes the client's spreadsheet from clinics.json.

    python scripts/nppes/mn-specialists/xlsx.py [--pulled YYYY-MM-DD] [--out PATH]

Needs openpyxl (pip install openpyxl). The Summary sheet counts with COUNTIFS
formulas over the All Clinics sheet, so it stays right if the client sorts,
edits or deletes rows. openpyxl does not compute formulas: open the file once in
Excel (or run it through LibreOffice) before sending it, so the counts are
cached for previewers that do not calculate.
"""

import argparse
import datetime
import json
from collections import Counter
from pathlib import Path
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "data" / "nppes" / "mn-specialists"

GROUPS = ["Injury Physicians", "Chiropractic", "Physical Therapy", "Imaging", "TMJ Specialists", "Oral Surgery"]
# Added 2026-10-19 from scripts/nppes/mn-tmj/, and checked on the web row by row.
TMJ, OMS = "TMJ Specialists", "Oral Surgery"
DENTAL = [TMJ, OMS]
# build.ts marks a row `recent` when its registry entry was updated since this year.
RECENT_YEAR = 2022
GROUP_NOTE = {
    "Injury Physicians": "Orthopedics, spine, sports medicine, neurosurgery, pain, PM&R, neurology",
    "Chiropractic": "Chiropractic clinics",
    "Physical Therapy": "Physical therapy clinics",
    "Imaging": "Radiology and MRI centers",
    TMJ: "TMJ and orofacial pain specialists and clinics",
    OMS: "Oral and maxillofacial surgeons (TMJ and jaw surgery)",
}
REGIONS = [
    "Twin Cities Metro",
    "Central Minnesota",
    "Northeast Minnesota",
    "Northwest Minnesota",
    "West Central Minnesota",
    "South Central Minnesota",
    "Southeast Minnesota",
    "Southwest Minnesota",
]

# (header, width, value-from-row)
COLUMNS = [
    ("Clinic Name", 46, lambda r: r["name"]),
    ("Specialty Group", 18, lambda r: r["group"]),
    ("Also In", 24, lambda r: ", ".join(g for g in r["groups"] if g != r["group"])),
    ("Specialties", 34, lambda r: ", ".join(r["specialties"])),
    ("Street", 30, lambda r: r["street"]),
    ("Suite", 12, lambda r: r["suite"]),
    ("City", 18, lambda r: r["city"]),
    ("County", 15, lambda r: r["county"]),
    ("Region", 22, lambda r: r["region"]),
    ("ZIP", 8, lambda r: r["zip"]),
    ("Phone", 15, lambda r: r["phone"]),
    ("Providers at Address", 11, lambda r: r["providers"]),
    ("On Xpert Connect", 11, lambda r: "Yes" if r["onPlatform"] else "No"),
    ("Registry Updated", 11, lambda r: registry_date(r["updated"])),
    ("Legal Name", 46, lambda r: r["legalName"]),
    # A practice found on the web, not in the registry, has none.
    ("NPI", 12, lambda r: r["npis"][0] if r["npis"] else None),
    ("Map", 8, None),
    ("NPPES Record", 10, None),
]


def letters(columns):
    return {name: get_column_letter(i + 1) for i, (name, _, _) in enumerate(columns)}


COL = letters(COLUMNS)


def columns_for(group):
    """
    The TMJ and Oral Surgery sheets carry what the web check found: who
    practises there, the website, whether the clinic was seen open, and for a
    surgeon whether the practice treats TMJ. The other sheets were never
    checked row by row, and an empty Website column on 2,900 rows would read as
    missing data. Every other sheet keeps COLUMNS, so COL is right for them.
    """
    if group not in DENTAL:
        return COLUMNS
    cols = list(COLUMNS)
    names = (
        ("TMJ Specialists Here", 34, lambda r: "; ".join(r["orofacial"]))
        if group == TMJ
        else ("Oral Surgeons Here", 34, lambda r: "; ".join(r["surgeons"]))
    )
    cols.insert(cols.index(next(c for c in cols if c[0] == "Specialties")) + 1, names)
    web = [
        ("Website", 26, None),
        ("Web Check", 12, lambda r: WEB_CHECK.get(r["webCheck"], "Not checked")),
    ]
    if group == OMS:
        web.append(("Treats TMJ", 9, lambda r: {True: "Yes", False: "No"}.get(r["treatsTmj"], "")))
    at = cols.index(next(c for c in cols if c[0] == "Phone")) + 1
    cols[at:at] = web
    return cols


# What the web check found, as the sheet says it.
WEB_CHECK = {
    "Confirmed": "Confirmed",
    "Moved": "New address",
    "Unverified": "Unverified",
}

FONT = "Arial"
NAVY = "1F3A5F"
TEXT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TITLE = Font(name=FONT, size=18, bold=True, color=NAVY)
SUBTITLE = Font(name=FONT, size=10, italic=True, color="595959")
BIG = Font(name=FONT, size=28, bold=True, color=NAVY)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
YES_FILL = PatternFill("solid", fgColor="E2EFDA")


def registry_date(iso):
    """The date the clinic last updated its NPI record, as a real date cell."""
    if not iso:
        return None
    y, m, d = (int(x) for x in iso.split("-"))
    return datetime.date(y, m, d)


def maps_url(r):
    q = f"{r['name']}, {r['street']}, {r['city']}, MN {r['zip']}"
    return "https://www.google.com/maps/search/?api=1&query=" + quote_plus(q)


def website_label(url):
    """'https://www.mnheadneckpain.com/locations/plymouth/' -> 'mnheadneckpain.com'."""
    host = url.split("//", 1)[-1].split("/", 1)[0]
    return host[4:] if host.startswith("www.") else host


def clinic_sheet(wb, title, rows, table_name, group=None):
    ws = wb.create_sheet(title)
    columns = columns_for(group)
    col = letters(columns)
    for i, (header, width, _) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=header)
        c.font = HEAD
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30

    for n, r in enumerate(rows, start=2):
        for i, (header, _, get) in enumerate(columns, start=1):
            c = ws.cell(row=n, column=i)
            if header == "Map":
                c.value = "Map"
                c.hyperlink = maps_url(r)
                c.font = LINK
            elif header == "NPPES Record":
                if r["npis"]:
                    c.value = "Record"
                    c.hyperlink = f"https://npiregistry.cms.hhs.gov/provider-view/{r['npis'][0]}"
                    c.font = LINK
                else:
                    c.value = "Web"
                    c.font = TEXT
            elif header == "Website":
                if r["website"]:
                    c.value = website_label(r["website"])
                    c.hyperlink = r["website"]
                    c.font = LINK
            else:
                c.value = get(r)
                c.font = TEXT
            if header in ("ZIP", "Phone", "NPI"):
                c.number_format = "@"
            if header == "Registry Updated":
                c.number_format = "mmm yyyy"
            if header in ("Providers at Address", "On Xpert Connect", "Registry Updated", "ZIP", "Map", "NPPES Record", "Web Check", "Treats TMJ"):
                c.alignment = Alignment(horizontal="center")

    last = len(rows) + 1
    ref = f"A1:{get_column_letter(len(columns))}{last}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight9", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)
    on = col["On Xpert Connect"]
    ws.conditional_formatting.add(
        f"{on}2:{on}{last}",
        FormulaRule(formula=[f'${on}2="Yes"'], fill=YES_FILL),
    )
    if "Web Check" in col:
        wc = col["Web Check"]
        ws.conditional_formatting.add(
            f"{wc}2:{wc}{last}",
            FormulaRule(formula=[f'${wc}2="Confirmed"'], fill=YES_FILL),
        )
        ws.conditional_formatting.add(
            f"{wc}2:{wc}{last}",
            FormulaRule(formula=[f'${wc}2="New address"'], fill=PatternFill("solid", fgColor="FFF2CC")),
        )
    ws.freeze_panes = "B2"
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    return ws


def summary_sheet(wb, rows, counties, pulled, tmj_pulled):
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    # County, Region, the groups, Total, Older entries: the widest table.
    for i in range(2, len(GROUPS) + 5):
        ws.column_dimensions[get_column_letter(i)].width = 13

    all_ = "'All Clinics'"
    grp = f"{all_}!${COL['Specialty Group']}:${COL['Specialty Group']}"
    reg = f"{all_}!${COL['Region']}:${COL['Region']}"
    cty = f"{all_}!${COL['County']}:${COL['County']}"
    onx = f"{all_}!${COL['On Xpert Connect']}:${COL['On Xpert Connect']}"
    name = f"{all_}!$A:$A"

    ws["A1"] = "Minnesota Specialist Clinics"
    ws["A1"].font = TITLE
    ws["A2"] = (
        f"Every organization registered in the federal NPI Registry (CMS) under these specialties "
        f"with a Minnesota practice address, its own name and a phone number. Pulled {pulled}. "
        f"TMJ specialists and oral surgeons added {tmj_pulled}, each one checked on the web."
    )
    ws["A2"].font = SUBTITLE
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:J2")
    ws.row_dimensions[2].height = 28

    ws["A4"] = f"=COUNTA({name})-1"
    ws["A4"].font = BIG
    ws["A4"].number_format = "#,##0"
    ws["A4"].alignment = Alignment(horizontal="left")
    ws["A5"] = "clinic locations across Minnesota"
    ws["A5"].font = BOLD
    ws["C4"] = f'=COUNTIF({onx},"No")'
    ws["C4"].font = BIG
    ws["C4"].number_format = "#,##0"
    ws["C4"].alignment = Alignment(horizontal="left")
    ws.merge_cells("C4:E4")
    ws["C5"] = "of them are not on Xpert Connect yet"
    ws["C5"].font = BOLD
    ws["A6"] = (
        "=\"Plus \"&TEXT(COUNTA('Older Entries'!$A:$A)-1,\"#,##0\")&\" older registry entries on the "
        f"'Older Entries' sheet: clinics that have not updated their federal registration since before {RECENT_YEAR}. "
        "Call to confirm before using them.\""
    )
    ws["A6"].font = Font(name=FONT, size=10, italic=True, color="9C5700")
    ws["A6"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A6:J6")
    ws.row_dimensions[6].height = 28

    def header(row, labels):
        for i, label in enumerate(labels, start=1):
            c = ws.cell(row=row, column=i, value=label)
            c.font = HEAD
            c.fill = HEAD_FILL
            c.border = BOX
            c.alignment = Alignment(horizontal="left" if i == 1 else "center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    def cell(row, col, value, fmt="#,##0", font=TEXT):
        c = ws.cell(row=row, column=col, value=value)
        c.font = font
        c.border = BOX
        if col > 1:
            c.number_format = fmt
            c.alignment = Alignment(horizontal="center")
        return c

    # By group. A clinic appears in every group sheet it belongs to, so these
    # count the group sheets, and their sum can exceed the total.
    r = 8
    ws.cell(row=r, column=1, value="By specialty group").font = BOLD
    r += 1
    header(r, ["Group", "Clinics", "Not on Xpert Connect", "Older entries", "What it includes"])
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=10)
    main_col = f"'Older Entries'!${COL['Specialty Group']}:${COL['Specialty Group']}"
    also_col = f"'Older Entries'!${COL['Also In']}:${COL['Also In']}"
    for g in GROUPS:
        r += 1
        sheet = f"'{g}'"
        on = letters(columns_for(g))["On Xpert Connect"]
        cell(r, 1, g, font=BOLD)
        cell(r, 2, f"=COUNTA({sheet}!$A:$A)-1")
        cell(r, 3, f'=COUNTIF({sheet}!${on}:${on},"No")')
        # Same membership rule as the group sheets: main group, or listed under Also In.
        cell(r, 4, f'=COUNTIF({main_col},"{g}")+COUNTIF({also_col},"*{g}*")')
        note = ws.cell(row=r, column=5, value=GROUP_NOTE[g])
        note.font = Font(name=FONT, size=9, color="595959")
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=10)
    r += 1
    foot = ws.cell(
        row=r,
        column=1,
        value="A multi-specialty clinic is listed in each of its groups, so the groups add up to more than the total.",
    )
    foot.font = Font(name=FONT, size=9, italic=True, color="595959")

    # Region x main group.
    r += 2
    ws.cell(row=r, column=1, value="By region (each clinic counted once, under its main specialty group)").font = BOLD
    r += 1
    header(r, ["Region"] + GROUPS + ["Total"])
    first = r + 1
    last_group = get_column_letter(len(GROUPS) + 1)
    for region in REGIONS:
        r += 1
        cell(r, 1, region, font=BOLD)
        for i, g in enumerate(GROUPS, start=2):
            cell(r, i, f'=COUNTIFS({reg},$A{r},{grp},"{g}")')
        cell(r, len(GROUPS) + 2, f"=SUM(B{r}:{last_group}{r})", font=BOLD)
    r += 1
    cell(r, 1, "Total", font=BOLD)
    for i in range(2, len(GROUPS) + 3):
        col = get_column_letter(i)
        cell(r, i, f"=SUM({col}{first}:{col}{r - 1})", font=BOLD)

    # County, all 87, most clinics first. Zeros stay visible.
    r += 2
    ws.cell(row=r, column=1, value="By county (all 87 Minnesota counties)").font = BOLD
    r += 1
    header(r, ["County", "Region"] + GROUPS + ["Total", "Older entries"])
    by_county = Counter(x["county"] for x in rows)
    for county, region in sorted(counties.items(), key=lambda kv: (-by_county.get(kv[0], 0), kv[0])):
        r += 1
        cell(r, 1, county, font=BOLD)
        c = cell(r, 2, region)
        c.alignment = Alignment(horizontal="left")
        c.font = Font(name=FONT, size=9, color="595959")
        for i, g in enumerate(GROUPS, start=3):
            cell(r, i, f'=COUNTIFS({cty},$A{r},{grp},"{g}")')
        cell(r, len(GROUPS) + 3, f"=SUM(C{r}:{get_column_letter(len(GROUPS) + 2)}{r})", font=BOLD)
        cell(r, len(GROUPS) + 4, f"=COUNTIF('Older Entries'!${COL['County']}:${COL['County']},$A{r})")
    # The region column is wider than the numbers in the county table.
    ws.column_dimensions["B"].width = 22
    ws.freeze_panes = "A8"


def about_sheet(wb, rows, verification, pulled, tmj_pulled, tmj_summary):
    ws = wb.create_sheet("About")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Minnesota Specialist Clinics: how this list was made", TITLE),
        ("", None),
        ("Source", BOLD),
        (
            "NPPES, the National Plan and Provider Enumeration System run by CMS (npiregistry.cms.hhs.gov). "
            "Every healthcare provider that bills in the US is registered there. It is public-domain data.",
            TEXT,
        ),
        (f"Pulled {pulled} through the NPI Registry API, Minnesota ZIP code by ZIP code, so the whole state is covered.", TEXT),
        ("", None),
        ("What counts as a clinic here", BOLD),
        ("An organization (not an individual) with its own registered name, a street address in Minnesota and a phone number.", TEXT),
        ("One row per clinic location. A group with several locations appears once per location.", TEXT),
        ("", None),
        ("Main list and Older Entries", BOLD),
        (
            f"The main sheets hold clinics whose registry entry was updated in {RECENT_YEAR} or later. "
            "Clinics are required to keep it current, so a recent update means the clinic itself reviewed its entry in the last few years.",
            TEXT,
        ),
        (
            f"The Older Entries sheet holds clinics that have not touched their entry since before {RECENT_YEAR}. Many are still open, "
            "but a web check found a large share had moved or closed, so call before relying on one.",
            TEXT,
        ),
        ("", None),
        ("Specialty groups", BOLD),
    ]
    for g in GROUPS:
        lines.append((f"{g}: {GROUP_NOTE[g]}.", TEXT))
    lines += [
        ("", None),
        (f"TMJ Specialists and Oral Surgery (added {tmj_pulled})", BOLD),
        (
            "TMJ Specialists: dentists registered in Orofacial Pain, the dental specialty for TMJ disorders; clinics whose "
            "name says TMJ, TMD or facial pain; and dental practices found on the web that make TMJ a focus, in their name, "
            "as a headline service, or through orofacial-pain board certification or training. A specialist whose practice "
            "is not registered as a TMJ clinic is listed under the practice the web check found, or under their own name.",
            TEXT,
        ),
        (
            "To find TMJ practices the registry does not label as TMJ, the whole state was searched on the web: the Twin "
            "Cities and its suburbs, more than 50 Greater Minnesota cities and towns, the regional health systems "
            "(Essentia, Sanford, CentraCare, Mayo Clinic Health System, Gundersen and others), and the member directories "
            "of the American Board of Orofacial Pain, the American Academy of Orofacial Pain and the American Academy of "
            "Craniofacial Pain. Outside the Twin Cities, Rochester and St. Cloud, few dentists make TMJ a focus.",
            TEXT,
        ),
        (
            "Oral Surgery: oral and maxillofacial surgeons, who operate on the jaw joint and treat jaw and facial injuries. "
            "Treats TMJ says whether the practice's own website offers TMJ treatment.",
            TEXT,
        ),
        (
            f"Every clinic in these two groups was looked up on the web in {tmj_pulled.split(' ')[0]} {tmj_pulled.split(' ')[-1]}. "
            "Website is the clinic's own site. Web Check: Confirmed means it was seen operating at the address shown; "
            "New address means it moved and the address shown is the new one; Unverified means it could not be confirmed "
            "either way. Clinics found closed were removed, and clinics confirmed open are on the main list even when "
            "their registry entry is older.",
            TEXT,
        ),
        (
            "TMJ Specialists Here and Oral Surgeons Here name the individual specialists registered at that address with "
            "the clinic's phone number, whose own registry entry has been updated since "
            f"{RECENT_YEAR}.",
            TEXT,
        ),
        ("", None),
        ("Left out on purpose", BOLD),
        (
            "Home-care, in-home and mobile services; therapy provided inside nursing homes and senior communities; "
            "school districts; hospice; equipment suppliers; animal chiropractic; radiologist groups that only read scans "
            "and see no patients; addresses that are apartments or PO boxes; general dentists who only list TMJ among "
            "many other services.",
            TEXT,
        ),
        (
            "Branch clinics of large hospital systems are listed only at the system's main address, because the registry "
            "does not say which specialties each branch offers.",
            TEXT,
        ),
        ("", None),
        ("Columns", BOLD),
        ("Specialty Group: the clinic's main group. Also In: other groups it belongs to; it is listed on those sheets too.", TEXT),
        ("Providers at Address: individual providers of the same specialty registered at that building (0 = none registered there individually).", TEXT),
        ("Registry Updated: when the clinic last updated its federal registry entry.", TEXT),
        ("On Xpert Connect: Yes if the clinic, or its health system at the same address, is already in the Xpert Connect directory.", TEXT),
        ("Map opens the address in Google Maps. NPPES Record opens the clinic's entry in the federal registry.", TEXT),
        ("", None),
        ("Limits", BOLD),
        (
            "The registry has no website or email for clinics. Websites are included only for TMJ and Oral Surgery, "
            "where each clinic was looked up; there are no emails.",
            TEXT,
        ),
        (
            "Clinics update their own registry entries. A clinic that moved or closed may still show its old address until it updates.",
            TEXT,
        ),
    ]
    if verification or tmj_summary:
        lines += [("", None), ("Spot check", BOLD)]
    if verification:
        lines.append((verification, TEXT))
    if tmj_summary:
        lines.append((tmj_summary, TEXT))
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if font:
            c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
        # Excel does not grow a row to fit wrapped text it did not lay out
        # itself, so the height is set here: ~105 characters to a line.
        if font is TEXT:
            ws.row_dimensions[i].height = 14 * max(1, -(-len(text) // 105))
    ws.row_dimensions[1].height = 30


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pulled", default="September 24, 2026")
    ap.add_argument("--tmj-pulled", default="October 19, 2026")
    ap.add_argument("--out", default=str(ROOT / "docs" / "Minnesota Specialist Clinics - 2026-09.xlsx"))
    args = ap.parse_args()

    rows = json.loads((DATA / "clinics.json").read_text(encoding="utf-8"))
    counties = {}
    import re

    regions_js = (ROOT / "scripts" / "lib" / "mn-regions.js").read_text(encoding="utf-8")
    for county, region in re.findall(r"'([^']+)':\s*'([^']+)'", regions_js):
        counties[county] = region
    assert len(counties) == 87, f"expected 87 counties, got {len(counties)}"

    verification = None
    vpath = DATA / "verification.json"
    if vpath.exists():
        verification = json.loads(vpath.read_text(encoding="utf-8")).get("summary")

    recent = [r for r in rows if r["recent"]]
    older = [r for r in rows if not r["recent"]]

    # The TMJ web check, in one sentence. Closed clinics are not rows any more,
    # so they are counted from web.json.
    tmj_summary = None
    dental = [r for r in rows if any(g in DENTAL for g in r["groups"])]
    checked = [r for r in dental if r["webCheck"]]
    web_path = ROOT / "data" / "nppes" / "mn-tmj" / "web.json"
    if checked and web_path.exists():
        # Only the registry rows: a sleep practice that closed was never a TMJ row.
        candidates = json.loads((web_path.parent / "web" / "candidates.json").read_text(encoding="utf-8"))
        keys = {c["key"] for c in candidates}
        checks = json.loads(web_path.read_text(encoding="utf-8"))["checks"]
        closed = sum(1 for k, c in checks.items() if k in keys and c["status"] == "Closed")
        by = Counter(r["webCheck"] for r in checked)
        tmj_summary = (
            f"TMJ and oral surgery: {len(checked) + closed} clinics looked up on the web in {args.tmj_pulled.split(' ')[0]} "
            f"{args.tmj_pulled.split(' ')[-1]}. {by['Confirmed']} were confirmed at their address, {by['Moved']} had moved "
            f"(the new address is shown), {by['Unverified']} could not be confirmed, and {closed} had closed and were removed."
        )

    wb = Workbook()
    summary_sheet(wb, recent, counties, args.pulled, args.tmj_pulled)
    clinic_sheet(wb, "All Clinics", recent, "AllClinics")
    for g in GROUPS:
        clinic_sheet(wb, g, [r for r in recent if g in r["groups"]], g.replace(" ", ""), group=g)
    clinic_sheet(wb, "Older Entries", older, "OlderEntries")
    about_sheet(wb, rows, verification, args.pulled, args.tmj_pulled, tmj_summary)
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{len(recent)} clinics (+{len(older)} older entries) -> {out}")
    for g in GROUPS:
        print(f"  {g}: {sum(1 for r in recent if g in r['groups'])} (+{sum(1 for r in older if g in r['groups'])} older)")


if __name__ == "__main__":
    main()
